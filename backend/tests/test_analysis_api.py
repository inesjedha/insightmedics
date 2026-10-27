"""Chaîne Analyses + résultats de bout en bout par l'API admin, sans IA (plan écrit à la
main, texte non rédigé) : commande, base corrigée, contre-audit, plan, calculs,
livrables, livraison sur le lien privé."""

import io
from datetime import UTC, datetime, timedelta

import numpy as np
import pandas as pd
from fastapi.testclient import TestClient

from app.db import SessionLocal
from app.main import app
from app.models import Audit, utcnow_iso

client = TestClient(app)
RNG = np.random.default_rng(3)


def _audit(token: str) -> str:
    db = SessionLocal()
    audit_id = f"audit_{token}"
    db.query(Audit).filter(Audit.id == audit_id).delete()
    db.add(Audit(id=audit_id, file_name="these.xlsx", file_size=1, stored_path="/x",
                 started_at=utcnow_iso(), status="done", token=token, paid=True,
                 expires_at=(datetime.now(UTC) + timedelta(days=2)).isoformat(),
                 score=61.0, row_count=40, column_count=4, missing_pct=0.0,
                 duplicates_pct=0.0, issues=[], needs_human_review=False,
                 score_detail={"score_final": 61},
                 profiling={"structure": {"n_rows": 40}, "columns": []},
                 ai_audit={"assessment": {"checklist": [
                     {"patients": "3", "row_ids": ["3"], "colonne": "Volume", "valeur": "9999",
                      "check": "Valeur exacte ?"}]}}))
    db.commit()
    db.close()
    return audit_id


def _csv() -> bytes:
    df = pd.DataFrame({"Numéro": range(1, 41),
                       "Groupe": ["A", "B"] * 20,
                       "Volume": RNG.normal(600, 80, 40).round(),
                       "Succès": RNG.choice(["OUI", "NON"], 40, p=[0.7, 0.3])})
    return df.to_csv(index=False).encode()


PLAN = {"title": "Test", "derived": [
            {"name": "succes", "kind": "recode", "type": "categorical", "source": "Succès",
             "mapping": {"OUI": "Oui", "NON": "Non"}}],
        "groups": {"variable": "Groupe"},
        "analyses": [
            {"kind": "proportion", "id": "p", "title": "Succès", "variable": "succes",
             "event": "Oui", "event_label": "Succès"},
            {"kind": "compare", "id": "c", "title": "Volume selon le groupe",
             "variables": [{"name": "Volume", "label": "Volume (mL)", "type": "quantitative"}]}],
        "primary": "p"}


def test_chaine_complete_sans_ia(admin_headers):
    audit_id = _audit("tokanalyse")
    h = admin_headers
    assert client.post(f"/audit/{audit_id}/analysis").status_code == 401  # admin seulement
    a = client.post(f"/audit/{audit_id}/analysis", headers=h).json()
    assert a["status"] == "awaiting_base" and a["auditToken"] == "tokanalyse"
    again = client.post(f"/audit/{audit_id}/analysis", headers=h).json()
    assert again["id"] == a["id"]                                  # idempotent
    aid = a["id"]

    # Base corrigée par l'étudiant (facultative) : elle remplace celle de l'audit
    r = client.post(f"/analyses/{aid}/base", headers=h,
                    files={"file": ("corrigee.csv", io.BytesIO(_csv()), "text/csv")})
    assert r.status_code == 200 and r.json()["fileName"] == "corrigee.csv"
    client.put(f"/analyses/{aid}/instructions", headers=h,
               json={"text": "Succès = OUI dans la colonne Succès (validé par le directeur)."})

    # Plan IA : sans clé, contre-audit fait et plan à écrire à la main
    r = client.post(f"/analyses/{aid}/plan/generate", headers=h)
    assert r.status_code == 202
    a = client.get(f"/analyses/{aid}", headers=h).json()
    assert a["status"] == "awaiting_base" and "à la main" in a["error"]
    assert a["recheck"]["checked_values"][0]["statut"] == "modifiée"   # 9999 corrigé

    bad = client.put(f"/analyses/{aid}/plan", headers=h, json={**PLAN, "primary": "zz"})
    assert bad.status_code == 409 and "absente" in bad.json()["detail"]
    a = client.put(f"/analyses/{aid}/plan", headers=h, json=PLAN).json()
    assert a["status"] == "plan_ready"

    # Livraison impossible avant relecture ; calculs (texte non rédigé sans IA)
    assert client.post(f"/analyses/{aid}/deliver", headers=h).status_code == 409
    assert client.post(f"/analyses/{aid}/compute", headers=h).status_code == 202
    a = client.get(f"/analyses/{aid}", headers=h).json()
    assert a["status"] == "review" and a["results"][0]["id"] == "p"
    assert a["flow"][-1]["n"] == 40
    for kind in ("resultats.docx", "resultats.xlsx"):
        f = client.get(f"/analyses/{aid}/{kind}", headers=h)
        assert f.status_code == 200 and len(f.content) > 1000, kind

    # Le client ne voit rien avant la livraison, puis ses deux fichiers
    assert client.get("/r/tokanalyse/analysis").status_code == 404
    assert client.get("/r/tokanalyse/analysis/resultats.docx").status_code == 404
    a = client.post(f"/analyses/{aid}/deliver", headers=h).json()
    assert a["status"] == "delivered"
    body = client.get("/r/tokanalyse/analysis").json()
    assert body["files"] == ["resultats.docx", "resultats.xlsx"]
    assert client.get("/r/tokanalyse/analysis/resultats.xlsx").status_code == 200
    assert client.get("/r/tokanalyse/analysis/autre.txt").status_code == 404
    listing = client.get("/analyses", headers=h).json()
    assert any(x["id"] == aid and x["auditFileName"] == "these.xlsx" for x in listing)


def test_commande_refusee_sur_un_apercu(admin_headers):
    db = SessionLocal()
    db.query(Audit).filter(Audit.id == "audit_tokapercu2").delete()
    db.add(Audit(id="audit_tokapercu2", file_name="b.xlsx", file_size=1, stored_path="/x",
                 started_at=utcnow_iso(), status="preview", token=None, paid=False))
    db.commit()
    db.close()
    r = client.post("/audit/audit_tokapercu2/analysis", headers=admin_headers)
    assert r.status_code == 409


def test_nettoyage_par_nos_soins_sur_la_base_de_l_audit(admin_headers, tmp_path):
    """Sans base corrigée : la base de l'audit est nettoyée par les corrections du plan,
    chaque cellule modifiée est consignée, et le classeur livre le journal."""
    import openpyxl

    src = tmp_path / "these.csv"
    pd.DataFrame({"N°": range(1, 9),
                  "Âge": [34, 999, 51, -3, 45, 60, 29, 999],
                  "Sexe": ["H", "h ", "Homme", "F", "femme", "F", "H", "F"],
                  "Volume": ["600 mL", "1 200 mL", "> 400", "450", "700", "800 ml", "?", "650"],
                  }).to_csv(src, index=False)
    audit_id = _audit("toknettoyage")
    db = SessionLocal()
    audit = db.get(Audit, audit_id)
    audit.stored_path, audit.file_name = str(src), "these.csv"
    db.commit()
    db.close()
    h = admin_headers
    aid = client.post(f"/audit/{audit_id}/analysis", headers=h).json()["id"]
    plan = {"title": "Nettoyage", "corrections": [
                {"kind": "missing_codes", "column": "Âge", "values": [999],
                 "reason": "999 = âge non renseigné"},
                {"kind": "set_missing", "column": "Âge", "where": 'col("Âge") < 0',
                 "reason": "Âge négatif impossible"},
                {"kind": "harmonize", "column": "Sexe", "reason": "Même réponse, écrite "
                 "différemment", "mapping": {"h": "Homme", "homme": "Homme", "f": "Femme",
                                             "femme": "Femme"}},
                {"kind": "to_number", "column": "Volume", "reason": "Volumes saisis avec "
                 "leur unité"},
                {"kind": "set_value", "column": "Volume", "lines": [9], "value": 640,
                 "source": "dossier vérifié par l'étudiant", "reason": "Valeur relevée"}],
            "to_check": [{"lines": [3], "column": "Volume", "question": "1 200 mL ?"}],
            "analyses": [{"kind": "describe", "id": "d", "title": "Patients", "by_group": False,
                          "variables": [{"name": "Âge", "type": "quantitative"},
                                        {"name": "Sexe", "type": "categorical"},
                                        {"name": "Volume", "type": "quantitative"}]}]}
    a = client.put(f"/analyses/{aid}/plan", headers=h, json=plan).json()
    assert a["status"] == "plan_ready", a
    assert client.post(f"/analyses/{aid}/compute", headers=h).status_code == 202
    a = client.get(f"/analyses/{aid}", headers=h).json()
    journal = {d["kind"]: d for d in a["derivedLog"] if d.get("step") == "correction"}
    assert journal["missing_codes"]["n"] == 2 and journal["missing_codes"]["lines"] == [3, 9]
    assert journal["set_missing"]["lines"] == [5]
    assert journal["harmonize"]["n"] == 7                          # « Homme » inchangé
    assert journal["to_number"]["illisibles"] == [4, 8]           # « > 400 », « ? »
    assert journal["set_value"]["lines"] == [9] and journal["set_value"]["source"]
    table = {r[0]: r[1] for r in a["results"][0]["table"]["rows"]}
    assert any("Homme" in k for k in table) and not any(k.strip() == "h" for k in table)

    wb = openpyxl.load_workbook(io.BytesIO(
        client.get(f"/analyses/{aid}/resultats.xlsx", headers=h).content))
    assert "Journal des corrections" in wb.sheetnames
    text = " ".join(str(c.value) for row in wb["Journal des corrections"].iter_rows()
                    for c in row if c.value is not None)
    assert "999 = âge non renseigné" in text and "« 1 200 mL » → 1200" in text
