"""Tests du classeur Excel remis à l'étudiant (4 onglets) — sans appel API."""

import io

from openpyxl import load_workbook
from tests.test_report_xlsx import SYNTH_AI, _profiling_min
from tests.test_score_engine import CLEAN, FULL_SI

from app.pipeline.report_client_xlsx import build_client_workbook
from app.pipeline.score_engine import compute_score

NEW_FIELDS = {
    "headline": {"title": "Exploitable après 3 vérifications", "subtitle": "Trois points bloquent."},
    "blocking_points": [{"title": "14 patients saisis deux fois", "what": "Inclusions 34 à 47.",
                         "why": "Double comptage.", "todo": "Vérifier les dossiers.",
                         "refs": ["4 ↔ 34"]}],
}


def _wb(assessment_extra=None, profiling=None):
    ai = {**SYNTH_AI, "assessment": {**SYNTH_AI["assessment"], **(assessment_extra or {})}}
    data = build_client_workbook(profiling or _profiling_min(), compute_score(CLEAN, FULL_SI), ai)
    return load_workbook(io.BytesIO(data))


def _values(ws):
    return [v for row in ws.iter_rows(values_only=True) for v in row if v not in (None, "")]


def _text(ws) -> str:
    return " ".join(str(v) for v in _values(ws))


def test_quatre_onglets_dans_l_ordre():
    wb = _wb(NEW_FIELDS)
    assert wb.sheetnames == ["Votre bilan", "Aux archives", "Questions au directeur",
                             "Pour aller plus loin"]
    assert wb.active.title == "Votre bilan"


def test_aux_archives_depuis_la_checklist_ou_les_anomalies():
    with_list = _wb({**NEW_FIELDS, "checklist": [{"patients": "4 et 34",
                                                  "check": "Même patient ou deux patients ?"}]})
    assert "n° 4 et 34" in _values(with_list["Aux archives"])
    old = _wb()  # SYNTH_AI : F002 demande une vérification au dossier
    assert "Modalité rare" in _text(old["Aux archives"])


def test_score_en_clair_sans_motifs_internes():
    vals = _text(_wb(NEW_FIELDS)["Pour aller plus loin"])
    assert "Votre note, point par point" in vals
    for token in ("statut :", "M5", "hash"):
        assert token not in vals


def test_bilan_verdict_et_points_bloquants():
    ws = _wb(NEW_FIELDS)["Votre bilan"]
    assert "Exploitable après 3 vérifications" in str(ws["A2"].value)   # pastille + verdict
    vals = _text(ws)
    assert "14 patients saisis deux fois" in vals and "Double comptage." in vals


def test_ancien_audit_sans_nouveaux_champs_reste_lisible():
    assert "exploitable avec réserves" in str(_wb()["Votre bilan"]["A2"].value)


def test_points_releves_tries_par_priorite():
    vals = _values(_wb(NEW_FIELDS)["Pour aller plus loin"])
    first, last = vals.index("Avant toute analyse"), vals.index("Facultative")
    assert first < last


def test_identifiants_masques_si_la_colonne_identifiant_est_identifiante():
    prof = _profiling_min()
    prof["structure"] = {**prof["structure"], "id_column": "dossier"}
    prof["pii_candidates"] = [{"column": "dossier"}]
    findings = [{"id": "F1", "severity": "critique", "anomaly_class": "A",
                 "certainty": "certain", "row_ids": ["DOS-123"], "title_fr": "t",
                 "explanation_fr": "e", "requires_source_verification": True}]
    wb = _wb({"findings": findings}, prof)
    for name in wb.sheetnames:
        assert "DOS-123" not in _text(wb[name])


def test_questions_en_blocs_avec_decision_retenue():
    fields = {**NEW_FIELDS, "client_decisions": [{"question_fr": "Exclure ?", "choices": [
        {"label": "Les garder", "consequence": "Limites"},
        {"label": "Les exclure", "consequence": "Flux"}], "recommendation_fr": "Garder."}]}
    ws = _wb(fields)["Questions au directeur"]
    vals = _text(ws)
    assert "Option A" in vals and "Les exclure" in vals and "Flux" in vals
    assert "Décision retenue" in vals and "Notre conseil" in vals
    assert ws.data_validations.dataValidation   # liste Option A / B / Autre


def test_bilan_v3_missions_avancement_et_statuts():
    fields = {**NEW_FIELDS,
              "missions": [{"feu": "rouge", "title": "Des patients saisis deux fois ?",
                            "vu": "Lignes 36 à 49.", "a_faire": ["Ouvrir les dossiers."],
                            "qui": "vous", "a_rouvrir": "28 dossiers"}],
              "deja_en_ordre": ["Chaque ligne a un numéro d'inclusion."],
              "checklist": [{"patients": "4 et 34", "lignes": "lignes 6 et 36",
                             "colonne": "toutes", "valeur": "identiques",
                             "check": "Même patient ?", "mission": 1}]}
    wb = _wb(fields)
    bilan = _text(wb["Votre bilan"])
    assert "Mission 1 · Rouge · Des patients saisis deux fois ?" in bilan
    assert "✓ Chaque ligne a un numéro d'inclusion." in bilan
    formulas = [c.value for row in wb["Votre bilan"].iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=COUNTIFS")]
    assert formulas and "Aux archives" in formulas[0]
    archives = wb["Aux archives"]
    vals = _values(archives)
    assert "lignes 6 et 36" in vals and "À vérifier" in vals and "Toute la ligne" in vals
    assert "Mission 1 · Rouge · Des patients saisis deux fois ?" in vals
    assert archives.conditional_formatting  # ligne verte une fois vérifiée


def test_impression_a4_et_pied_de_page():
    for ws in _wb(NEW_FIELDS).worksheets:
        assert str(ws.page_setup.paperSize) == str(ws.PAPERSIZE_A4)
        assert "Insight Medics" in ws.oddFooter.left.text


def test_textes_longs_jamais_coupes():
    long = "Une phrase très longue qui décrit un problème. " * 12
    ws = _wb({**NEW_FIELDS, "missions": [{"feu": "orange", "title": "Long", "vu": long}]})[
        "Votre bilan"]
    rows = [c.row for row in ws.iter_rows() for c in row if c.value == long]
    assert rows and ws.row_dimensions[rows[0]].height >= 15 * 4
