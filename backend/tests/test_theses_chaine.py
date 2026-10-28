"""Format faculté de bout en bout, avec une IA simulée : fiche de mission (admin et
étudiant), « Tout lancer » (plan → calculs → rédaction enchaînés), refus d'un brouillon
qui invente un chiffre, livrables Word à la charte des quatre facultés, contrôles et
livraison."""

import io
import json
from datetime import UTC, datetime, timedelta

import numpy as np
import pandas as pd
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.config import settings
from app.db import SessionLocal
from app.main import app
from app.models import Audit, utcnow_iso
from app.pipeline import llm_client

client = TestClient(app)
TOKEN = "tokthesefac"

PLAN = {
    "title": "Hémothorax",
    "derived": [{"name": "succes", "kind": "recode", "type": "binary", "source": "Succès",
                 "mapping": {"OUI": "Oui", "NON": "Non"}}],
    "groups": {"variable": "Groupe"},
    "analyses": [
        {"kind": "describe", "id": "t1", "title": "Patients", "by_group": False,
         "variables": [{"name": "Âge", "label": "Âge (ans)", "type": "quantitative"},
                       {"name": "Classe", "label": "Classe", "type": "categorical"}]},
        {"kind": "proportion", "id": "p", "title": "Succès", "variable": "succes",
         "event": "Oui", "event_label": "Succès"},
        {"kind": "compare", "id": "c", "title": "Selon le groupe",
         "variables": [{"name": "Volume", "label": "Volume (mL)", "type": "quantitative"},
                       {"name": "Diabète", "label": "Diabète", "type": "binary"}]},
        {"kind": "paired", "id": "s", "title": "SpO2", "label": "SpO2 (%)",
         "times": ["J0", "J1"], "variables": ["SpO2 J0", "SpO2 J1"]},
    ],
    "primary": "p",
}


def _draft(patients: int) -> dict:
    h = lambda level, text: {"type": "heading", "level": level, "text": text}  # noqa: E731
    p = lambda text: {"type": "paragraph", "text": text}  # noqa: E731
    return {
        "methods": [
            h(1, "Logiciel et seuil"),
            p("Les données ont été analysées avec [À CONFIRMER : logiciel d'analyse]. Le seuil "
              "de signification a été fixé à p < 0,05 pour des tests bilatéraux."),
            h(1, "Plan de l'analyse"),
            p("Le plan de l'analyse est résumé dans le {{TAB:plan}}."),
            {"type": "text_table", "key": "plan", "title": "Plan de l'analyse statistique",
             "columns": ["Partie des résultats", "Question posée", "Outils statistiques"],
             "rows": [["Analyse descriptive", "Qui sont les patients ?", "Effectifs, médianes"],
                      ["Critère principal", "Quel est le taux de succès ?", "IC de Wilson"]]},
            h(1, "Analyse descriptive"),
            p("Les variables qualitatives ont été exprimées en effectifs et en pourcentages."),
        ],
        "results": [
            h(0, "Analyse descriptive"),
            h(1, "Population de l'étude"),
            p(f"Au total, {patients} patients ont été analysés ({{{{FIG:flux}}}})."),
            {"type": "figure", "figure_id": "flux", "title": "Diagramme de flux de l'étude"},
            h(1, "Caractéristiques des patients"),
            p("Les caractéristiques des patients sont présentées dans le {{TAB:t1}} ; la "
              "répartition par classe est présentée dans la {{FIG:t1::Classe}}."),
            {"type": "table", "analysis_id": "t1", "title": "Caractéristiques des patients"},
            {"type": "figure", "figure_id": "t1::Classe", "title": "Répartition par classe"},
            h(0, "Critère de jugement principal"),
            p("Le taux de succès est présenté dans le {{TAB:p}}."),
            {"type": "table", "analysis_id": "p", "title": "Taux de succès"},
            h(0, "Comparaison des groupes"),
            p("Les deux groupes sont comparés dans le {{TAB:c}}."),
            {"type": "table", "analysis_id": "c", "title": "Comparaison selon le groupe"},
            h(0, "Évolution de la SpO2"),
            p("L'évolution de la SpO2 est présentée dans le {{TAB:s}} et la {{FIG:s}}."),
            {"type": "table", "analysis_id": "s", "title": "SpO2 avant et après"},
            {"type": "figure", "figure_id": "s", "title": "Évolution de la SpO2"},
        ],
        "to_confirm": ["logiciel d'analyse"],
        "missing": [],
    }


def _audit(tmp_path, token: str = TOKEN) -> str:
    rng = np.random.default_rng(11)
    n = 40
    src = tmp_path / "these.csv"
    pd.DataFrame({
        "Numéro": range(1, n + 1),
        "Groupe": ["Ponction", "Drain"] * (n // 2),
        "Âge": rng.integers(20, 76, n),
        "Classe": rng.choice(["Stade 1", "Stade 2", "Stade 3"], n),
        "Volume": rng.normal(600, 80, n).round(),
        "Diabète": rng.choice(["Oui", "Non"], n, p=[0.3, 0.7]),
        "Succès": rng.choice(["OUI", "NON"], n, p=[0.7, 0.3]),
        "SpO2 J0": rng.normal(93, 2, n).round(),
        "SpO2 J1": rng.normal(96, 2, n).round(),
    }).to_csv(src, index=False)
    db = SessionLocal()
    audit_id = "audit_" + token
    db.query(Audit).filter(Audit.id == audit_id).delete()
    db.add(Audit(id=audit_id, file_name="these.csv", file_size=1, stored_path=str(src),
                 started_at=utcnow_iso(), status="done", token=token, paid=True,
                 expires_at=(datetime.now(UTC) + timedelta(days=2)).isoformat(),
                 score=70.0, row_count=n, column_count=9, missing_pct=0.0,
                 duplicates_pct=0.0, issues=[], needs_human_review=False,
                 score_detail={"score_final": 70},
                 profiling={"structure": {"n_rows": n}, "columns": []},
                 ai_audit={"assessment": {"checklist": []}}))
    db.commit()
    db.close()
    return audit_id


@pytest.fixture()
def fake_claude(monkeypatch):
    """IA simulée : le plan, puis un brouillon qui invente un effectif (refusé), puis le
    bon brouillon."""
    calls = {"plan": 0, "draft": 0}

    def call_llm(system: str, prompt: str, effort=None) -> str:
        if "RÈGLES DE LA PLATEFORME" in system:
            calls["draft"] += 1
            assert "IX" in system or "FMS" in system                 # règles de la faculté
            assert "LOGICIEL À DÉCLARER" in prompt and "flux" in prompt
            return json.dumps(_draft(41 if calls["draft"] == 1 else 40))
        calls["plan"] += 1
        return json.dumps(PLAN)

    monkeypatch.setattr(llm_client, "call_llm", call_llm)
    monkeypatch.setattr(settings, "anthropic_api_key", "cle-factice")
    return calls


def _text(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs]
    parts += [c.text for t in doc.tables for row in t.rows for c in row.cells]
    # Espaces insécables rendues en espaces simples, pour des assertions lisibles.
    return "\n".join(parts).replace("\u00a0", " ")


def test_schema_public_de_la_fiche():
    body = client.get("/theses/fiche-schema").json()
    assert body["champs"]["these.faculte"]["valeurs"] == ["FMS", "FMT", "FMM", "FMSf"]


def test_format_faculte_de_bout_en_bout(admin_headers, fake_claude, tmp_path):
    h = admin_headers
    audit_id = _audit(tmp_path)
    aid = client.post(f"/audit/{audit_id}/analysis", headers=h).json()["id"]

    # Fiche : l'étudiant remplit ses champs, pas ceux de l'équipe
    r = client.put(f"/r/{TOKEN}/fiche", json={"these.titre": "Ponction pleurale "
                                              "évacuatrice échoguidée des hémothorax",
                                              "these.faculte": "FMS",
                                              "objectifs.principal": "évaluer l'efficacité",
                                              "these.periode": "je ne sais pas"})
    assert r.status_code == 200 and r.json()["editable"] is True
    team_only = client.put(f"/r/{TOKEN}/fiche", json={"multivarie.regle": "Entrée"})
    assert team_only.status_code == 422 and "équipe" in team_only.json()["detail"]
    bad = client.put(f"/analyses/{aid}/fiche", headers=h, json={"these.faculte": "Paris"})
    assert bad.status_code == 409 and "inconnue" in bad.json()["detail"]
    a = client.put(f"/analyses/{aid}/fiche", headers=h,
                   json={"population.unite": "patient"}).json()
    assert a["fiche"]["these.faculte"] == "FMS" and a["fiche"]["population.unite"] == "patient"

    # Un bouton : plan, calculs et rédaction enchaînés (premier brouillon refusé)
    r = client.post(f"/analyses/{aid}/launch", headers=h)
    assert r.status_code == 202
    a = client.get(f"/analyses/{aid}", headers=h).json()
    assert a["status"] == "review", a["error"]
    assert fake_claude == {"plan": 1, "draft": 2}
    assert a["draft"]["results"][2]["text"].startswith("Au total, 40 patients")
    assert a["deliverables"] == ["analyse_statistique.docx", "resultats.docx", "resultats.xlsx"]
    assert not [c for c in a["checks"] if c["level"] == "error"], a["checks"]
    assert any("À CONFIRMER" in c["message"] for c in a["checks"])
    assert client.put(f"/r/{TOKEN}/fiche", json={"these.titre": "x"}).status_code == 409

    # Livrables Word à la charte
    stats = client.get(f"/analyses/{aid}/analyse_statistique.docx", headers=h)
    assert stats.status_code == 200
    assert "Analyse_statistique_These_Ponction.docx" in stats.headers["content-disposition"]
    text = _text(stats.content)
    assert "IX. ANALYSE STATISTIQUE :" in text and "Tableau I :" in text
    assert "1. Logiciel et seuil :" in text
    results = client.get(f"/analyses/{aid}/resultats.docx", headers=h)
    text = _text(results.content)
    assert "I. ANALYSE DESCRIPTIVE :" in text
    # Tableaux du chapitre numérotés à la suite de celui du M&M ; flux en Figure 1
    assert "Tableau II : Caractéristiques des patients" in text
    assert "Figure 1 : Diagramme de flux" in text and "Figure 2" in text
    assert "Test" not in [c.text for c in Document(io.BytesIO(results.content)).tables[0]
                          .rows[0].cells]
    # « Classe » est en figure : retirée du tableau descriptif
    first = Document(io.BytesIO(results.content)).tables[0]
    assert not any(row.cells[0].text.startswith("Classe") for row in first.rows)
    assert len(Document(io.BytesIO(results.content)).inline_shapes) == 3

    # Mêmes contenus aux trois autres chartes
    for fac, heading in (("FMT", "6. Analyse statistique"), ("FMM", "7. Analyse statistique"),
                         ("FMSf", "8. ANALYSE STATISTIQUE")):
        client.put(f"/analyses/{aid}/fiche", headers=h, json={"these.faculte": fac})
        stats = client.get(f"/analyses/{aid}/analyse_statistique.docx", headers=h)
        assert heading in _text(stats.content), fac
        assert client.get(f"/analyses/{aid}/resultats.docx", headers=h).status_code == 200

    # Livraison : l'étudiant voit ses trois fichiers
    a = client.post(f"/analyses/{aid}/deliver", headers=h).json()
    assert a["status"] == "delivered"
    files = client.get(f"/r/{TOKEN}/analysis").json()["files"]
    assert files == ["analyse_statistique.docx", "resultats.docx", "resultats.xlsx"]
    assert client.get(f"/r/{TOKEN}/analysis/resultats.docx").status_code == 200


def test_lancement_refuse_si_la_fiche_est_incomplete(admin_headers, tmp_path):
    h = admin_headers
    audit_id = _audit(tmp_path, "tokficheincomplete")
    aid = client.post(f"/audit/{audit_id}/analysis", headers=h).json()["id"]
    r = client.post(f"/analyses/{aid}/launch", headers=h)
    assert r.status_code == 409 and "faculté" in r.json()["detail"]
    client.put(f"/analyses/{aid}/fiche", headers=h, json={"these.faculte": "FMT"})
    r = client.post(f"/analyses/{aid}/launch", headers=h)
    assert r.status_code == 409 and "Titre exact" in r.json()["detail"]
    client.put(f"/analyses/{aid}/fiche", headers=h,
               json={"these.titre": "T", "objectifs.principal": "O", "these.langue": "ar"})
    r = client.post(f"/analyses/{aid}/launch", headers=h)
    assert r.status_code == 409 and "arabe" in r.json()["detail"]
