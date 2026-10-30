"""Chartes des facultés : cohérence avec les JSON de Dr Hamza, typographie, fiche de
mission, rendu Word (numérotation, intercalaire, appels de tableaux)."""

import io
import json
from pathlib import Path

import pytest
from docx import Document

from app.theses import fiche
from app.theses.charte import CHARTES, FACULTIES, get_charte
from app.theses.document import (
    Chapter,
    HeadingBlock,
    ParagraphBlock,
    TableBlock,
    ThesisDoc,
)
from app.theses.render_docx import RenderError, render
from app.theses.typo import NBSP, fr_typo, roman

HAMZA = Path(__file__).parent.parent / "app" / "theses" / "hamza"


def _hamza(code: str) -> dict:
    return json.loads((HAMZA / f"04_Charte_{code}.json").read_text(encoding="utf-8"))


@pytest.mark.parametrize("code", FACULTIES)
def test_charte_identique_aux_mesures_de_hamza(code):
    h, c = _hamza(code), get_charte(code)
    m = h["page"]["marges_cm"]
    assert c.margins_cm == (m["haut"], m["bas"], m.get("gauche", m.get("gauche_reliure")),
                            m["droite"])
    assert c.usable_cm == h["page"]["largeur_utile_cm"]
    assert (c.body.name, c.body.size) == (h["texte"]["police"], h["texte"]["taille_pt"])
    assert c.body_para.line == h["texte"]["interligne"]
    assert c.body_para.first_indent_cm == h["texte"]["retrait_premiere_ligne_cm"]
    assert (c.table.font.name, c.table.font.size) == (h["tableaux"]["police"],
                                                      h["tableaux"]["taille_pt"])
    cap = h["legendes"]["tableau"]
    assert (c.caption_table.font.name, c.caption_table.font.size) == (cap["police"],
                                                                      cap["taille_pt"])
    assert c.refs.open == h["references"]["appel"][0]
    numbered = [t for t in h["titres"] if t.get("numerotation")]
    assert c.levels[0].suffix.strip() == numbered[0].get("fin", "").strip()
    assert c.roman_top_level == (numbered[0]["numerotation"] == "I.")


def test_typographie_francaise():
    t = fr_typo("p = 0,014 ; IC 95 % [2-36] : Tableau III et l'âge (n=30)")
    assert f"p{NBSP}={NBSP}0,014{NBSP};" in t
    assert f"IC{NBSP}95{NBSP}%" in t and "[2–36]" in t and f"Tableau{NBSP}III" in t
    assert f"n{NBSP}={NBSP}30" in t and "l’âge" in t
    assert fr_typo(t) == t                                   # idempotent
    assert "doi:10.1000/x:y" in fr_typo("voir doi:10.1000/x:y")  # DOI intact
    assert roman(9) == "IX" and roman(14) == "XIV"


def test_fiche_de_mission():
    clean = fiche.validate({"these.faculte": "FMSf", "population.effectif_annonce": 50,
                            "these.periode": "je ne sais pas", "these.titre": ""})
    assert clean == {"these.faculte": "FMSf", "population.effectif_annonce": 50,
                     "these.periode": "je ne sais pas"}
    with pytest.raises(fiche.FicheError) as err:
        fiche.validate({"inconnu": 1, "these.faculte": "Paris",
                        "population.effectif_annonce": "cinquante"})
    assert len(err.value.problems) == 3
    with pytest.raises(fiche.FicheError):
        fiche.validate({"multivarie.regle": "Entrée"}, student=True)
    assert "Titre exact de la fiche d'inscription" in fiche.missing_for_launch({})
    assert fiche.to_confirm(clean) == ["Période d'étude et date de point"]
    assert fiche.merge({"a": 1, "b": 2}, {"a": None, "c": 3}) == {"b": 2, "c": 3}


def _doc(code: str) -> bytes:
    blocks = [HeadingBlock(level=0, text="Analyse descriptive"),
              HeadingBlock(level=1, text="Population de l'étude"),
              ParagraphBlock(text="Les patients sont décrits dans le {{TAB:t}} (**p = 0,014**)."),
              TableBlock(key="t", title="Patients", columns=["Variable", "n (%)"],
                         rows=[["Sexe", ""], ["    Homme", "45 (90,0 %)"]])]
    return render(ThesisDoc(chapters=[Chapter(key="resultats", name="Résultats", blocks=blocks)],
                            first_table=3, first_page=18), CHARTES[code])


@pytest.mark.parametrize("code,first,second", [("FMS", "I. ANALYSE DESCRIPTIVE :",
                                                "1. Population de l’étude :"),
                                               ("FMT", "1. Analyse descriptive",
                                                "1.1. Population de l’étude"),
                                               ("FMSf", "1. ANALYSE DESCRIPTIVE",
                                                "1.1. Population de l’étude")])
def test_rendu_word_par_faculte(code, first, second):
    doc = Document(io.BytesIO(_doc(code)))
    texts = [p.text.replace(NBSP, " ") for p in doc.paragraphs]
    assert first in texts and second in texts
    assert any("dans le Tableau III" in t for t in texts)           # numéro de départ
    assert any(t == "Tableau III : Patients" for t in texts)
    bold = [r for p in doc.paragraphs for r in p.runs if r.bold and "0,014" in r.text]
    assert bold                                                     # p < 0,05 en gras
    assert doc.core_properties.author == ""
    assert doc.sections[-1].left_margin.cm == pytest.approx(CHARTES[code].margins_cm[2],
                                                            abs=0.01)


def test_appel_sans_tableau_refuse():
    blocks = [ParagraphBlock(text="Voir le {{TAB:absent}}.")]
    with pytest.raises(RenderError):
        render(ThesisDoc(chapters=[Chapter(key="resultats", name="Résultats", blocks=blocks)]),
               CHARTES["FMM"])


def test_boites_memes_quartiles_que_le_tableau():
    """La boîte d'une figure montre le « médiane [Q1–Q3] » du tableau (méthode SPSS)."""
    import pandas as pd

    from app.analysis.stats import quartiles
    from app.theses.figures import box_stats

    values = [35, 40, 50, 55, 55, 60, 65, 65, 70, 80, 80, 85, 95, 130]
    b = box_stats(values)
    assert (b["q1"], b["med"], b["q3"]) == quartiles(pd.Series(values))
    assert b["whishi"] == 95 and b["whislo"] == 35               # 130 hors des moustaches
