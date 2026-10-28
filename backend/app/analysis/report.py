"""Livrables de l'analyse : la partie Résultats (Word) et le classeur (Excel).

Word : paragraphe « Analyse statistique » à insérer dans le Matériel et méthodes,
population et diagramme de flux, puis une sous-partie par analyse (texte, tableau à trois
traits, figure pour les évolutions). Excel : les tableaux, la base d'analyse (colonnes
calculées, INCLUS_ANALYSE et motif, sans colonne identifiante), le diagramme de flux et
le journal des colonnes calculées.
"""

from __future__ import annotations

import datetime
import io
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Side

from ..pipeline import xlsx_style as st

BRAND = "1F3FD1"
MUTED = RGBColor(0x6B, 0x72, 0x80)


# ---------------------------------------------------------------- Word

def _border(cell: Any, **sides: str) -> None:
    """Trait horizontal d'une cellule (tableau à trois traits, style des thèses)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for side, size in sides.items():
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), "000000")
        borders.append(el)


def _table(doc: Any, columns: list[str], rows: list[list[Any]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(columns))
    table.autofit = True
    for j, title in enumerate(columns):
        cell = table.rows[0].cells[j]
        cell.text = ""
        run = cell.paragraphs[0].add_run(str(title))
        run.bold = True
        run.font.size = Pt(9.5)
        _border(cell, top="12", bottom="6")
    for i, row in enumerate(rows, start=1):
        for j, value in enumerate(row):
            cell = table.rows[i].cells[j]
            cell.text = ""
            text = str(value)
            run = cell.paragraphs[0].add_run(text.strip() if j == 0 else text)
            run.font.size = Pt(9.5)
            if j == 0 and text.startswith("    "):
                cell.paragraphs[0].paragraph_format.left_indent = Cm(0.4)
            elif j == 0 and len(row) > 1 and all(str(v) == "" for v in row[1:]):
                run.bold = True  # en-tête d'une variable qualitative
            if i == len(rows):
                _border(cell, bottom="12")


def _figure(fig_spec: dict[str, Any]) -> io.BytesIO:
    fig, ax = plt.subplots(figsize=(5.2, 2.8), dpi=200)
    ax.plot(fig_spec["times"], fig_spec["values"], marker="o", color=f"#{BRAND}")
    ax.set_ylabel(f"{fig_spec['label']} ({fig_spec['stat']})")
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="y", alpha=0.3)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png")
    plt.close(fig)
    buf.seek(0)
    return buf


def _note(doc: Any, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED


def build_results_docx(plan: dict[str, Any], flow: list[dict[str, Any]],
                       results: list[dict[str, Any]], text: dict[str, Any] | None) -> bytes:
    text = text or {}
    sections = {s["analysis_id"]: s for s in text.get("sections") or []}
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    doc.add_heading("Résultats", level=0)
    if plan.get("title"):
        _note(doc, plan["title"])

    if text.get("methods_paragraph"):
        doc.add_heading("Analyse statistique (à insérer dans le Matériel et méthodes)", level=2)
        doc.add_paragraph(text["methods_paragraph"])

    doc.add_heading("Population étudiée", level=1)
    if text.get("flow_paragraph"):
        doc.add_paragraph(text["flow_paragraph"])
    flow_rows = [[f["label"], str(f.get("excluded", "")),
                  str(f.get("remaining", f.get("n", "")))] for f in flow]
    _note(doc, "Diagramme de flux")
    _table(doc, ["Étape", "Lignes écartées", "Lignes restantes"], flow_rows)

    numbers = table_numbers(results)
    for r in results:
        k = numbers.get(r["id"])
        if k is None:
            continue
        s = sections.get(r["id"], {})
        doc.add_heading(s.get("heading") or r["title"], level=1)
        for para in s.get("paragraphs") or []:
            doc.add_paragraph(para)
        caption = doc.add_paragraph()
        cap = caption.add_run(s.get("table_title") or f"Tableau {k} : {r['title']}")
        cap.bold = True
        cap.font.size = Pt(10)
        _table(doc, r["table"]["columns"], r["table"]["rows"])
        for w in r.get("warnings") or []:
            _note(doc, f"Note : {w}")
        if r.get("figure"):
            doc.add_picture(_figure(r["figure"]), width=Cm(13))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            _note(doc, f"Figure : évolution de {r['figure']['label']} "
                       f"({r['figure']['stat']} par temps).")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- Excel

CORRECTION_LABEL = {
    "trim": "Espaces retirés",
    "missing_codes": "Code de case vide → vide",
    "harmonize": "Réponses écrites de la même façon",
    "to_number": "Texte → nombre",
    "to_date": "Texte → date",
    "set_missing": "Valeur impossible → vide",
    "set_value": "Valeur corrigée (dossier)",
}
KIND_LABEL = {"formula": "Calcul", "recode": "Réponses regroupées", "bins": "Classes",
              "parse_number": "Texte → nombre", "pct_change": "Variation en %"}
DOC = "Votre partie Résultats"
SHEET_SUMMARY, SHEET_TABLES = "Sommaire", "Tableaux"
SHEET_FLOW, SHEET_BASE = "Diagramme de flux", "Base d'analyse"
SHEET_DICT, SHEET_LOG, SHEET_DERIVED = ("Dictionnaire des colonnes", "Journal des corrections",
                                        "Colonnes calculées")

TOP_RULE = Border(top=Side(style="medium", color="111827"),
                  bottom=Side(style="thin", color="111827"))
BOTTOM_RULE = Border(bottom=Side(style="medium", color="111827"))


def table_numbers(results: list[dict[str, Any]]) -> dict[str, int]:
    """Numéro de chaque tableau affiché (les analyses en erreur ne prennent pas de
    numéro) : le même dans le Word et dans le classeur."""
    numbers: dict[str, int] = {}
    for r in results:
        if not r.get("error") and r["table"]["rows"]:
            numbers[r["id"]] = len(numbers) + 1
    return numbers


def _caption(r: dict[str, Any], k: int, sections: dict[str, dict[str, Any]]) -> str:
    return (sections.get(r["id"]) or {}).get("table_title") or f"Tableau {k} : {r['title']}"


def _three_line_table(ws: Any, row: int, caption: str, columns: list[str],
                      rows: list[list[Any]], notes: list[str] | None = None) -> int:
    """Tableau de thèse à trois traits (comme dans le Word) ; renvoie la ligne suivante."""
    st.merged_text(ws, row, 1, max(len(columns), 2), caption, st.font(11, bold=True))
    row += 1
    for j, c in enumerate(columns, start=1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.font = st.font(10.5, bold=True)
        cell.border = TOP_RULE
        cell.alignment = st.WRAP_CENTER if j > 1 else Alignment(vertical="center",
                                                                 wrap_text=True)
    st.fit_height(ws, row, [(c, st.col_width(ws, j), 10.5) for j, c in enumerate(columns, 1)],
                  minimum=20)
    for r in rows:
        row += 1
        for j, v in enumerate(r, start=1):
            cell = ws.cell(row=row, column=j, value=v)
            cell.font = st.BODY
            indented = j == 1 and isinstance(v, str) and v.startswith("  ")
            cell.alignment = Alignment(wrap_text=True, vertical="top",
                                       horizontal="left" if j == 1 else "center",
                                       indent=2 if indented else 0)
        st.fit_height(ws, row, [(v, st.col_width(ws, j), 11) for j, v in enumerate(r, 1)])
    for j in range(1, len(columns) + 1):
        ws.cell(row=row, column=j).border = BOTTOM_RULE
    row += 1
    for n in notes or []:
        row = st.merged_text(ws, row, 1, max(len(columns), 4), n,
                             st.font(9, italic=True, color=st.MUTED_INK))
    return row + 1


def _cell_value(v: Any) -> Any:
    if v is None or v is pd.NA or v is pd.NaT or (isinstance(v, float) and pd.isna(v)):
        return None
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    return v if isinstance(v, (int, float, str)) else str(v)


def _value_summary(s: pd.Series) -> str:
    counts = s.dropna().astype(str).str.strip().value_counts()
    if 0 < len(counts) <= 8:
        return ", ".join(f"{k} ({v})" for k, v in counts.items())
    num = pd.to_numeric(s, errors="coerce").dropna()
    if len(num) >= max(3, 0.5 * s.notna().sum()):
        return f"de {st.fr_num(float(num.min()))} à {st.fr_num(float(num.max()))}"
    return f"{len(counts)} réponses différentes" if len(counts) else "—"


def _kind(s: pd.Series) -> str:
    if pd.api.types.is_datetime64_any_dtype(s):
        return "Date"
    if pd.api.types.is_numeric_dtype(s):
        return "Nombre"
    return "Texte"


def build_results_xlsx(base: pd.DataFrame, pii_columns: set[str], flow: list[dict[str, Any]],
                       results: list[dict[str, Any]], derived_log: list[dict[str, Any]],
                       plan: dict[str, Any], text: dict[str, Any] | None = None) -> bytes:
    sections = {s["analysis_id"]: s for s in (text or {}).get("sections") or []}
    numbers = table_numbers(results)
    corrections = [d for d in derived_log if d.get("step") == "correction"]
    derived = [d for d in derived_log if d.get("step") != "correction"]
    wb = Workbook()
    summary = wb.active
    summary.title = SHEET_SUMMARY

    # --- Tableaux (trois traits, numérotés comme dans le Word)
    ws = wb.create_sheet(SHEET_TABLES)
    st.setup_sheet(ws, [44, 22, 22, 22, 22, 22], st.BRAND, DOC)
    row = st.title_block(ws, f"Insight Medics · {DOC}", "Les tableaux de vos résultats",
                         "Mêmes tableaux et même numérotation que dans le Word.", 6)
    anchors: list[tuple[str, int]] = []
    for r in results:
        k = numbers.get(r["id"])
        if k is None:
            continue
        caption = _caption(r, k, sections)
        anchors.append((caption, row))
        notes = [f"Note : {w}" for w in r.get("warnings") or []]
        if r.get("methods"):
            notes.append("Tests : " + " ; ".join(r["methods"]))
        row = _three_line_table(ws, row, caption, r["table"]["columns"], r["table"]["rows"],
                                notes)

    # --- Diagramme de flux
    ws = wb.create_sheet(SHEET_FLOW)
    st.setup_sheet(ws, [58, 18, 18], "11643A", DOC, landscape=False)
    row = st.title_block(ws, f"Insight Medics · {DOC}", "Diagramme de flux",
                         "Combien de lignes de votre base ont été analysées, et pourquoi les "
                         "autres ont été écartées (aucune n'est supprimée).", 3)
    _three_line_table(ws, row, "Population étudiée", ["Étape", "Lignes écartées",
                                                      "Lignes restantes"],
                      [[f["label"], f.get("excluded", ""), f.get("remaining", f.get("n", ""))]
                       for f in flow])

    # --- Base d'analyse : statut d'inclusion en tête, lignes écartées en gris
    ws = wb.create_sheet(SHEET_BASE)
    export = base.drop(columns=[c for c in base.columns if c in pii_columns])
    front = [c for c in ("INCLUS_ANALYSE", "MOTIF_EXCLUSION") if c in export.columns]
    export = export[front + [c for c in export.columns if c not in front]]
    widths = []
    for c in export.columns:
        sample = export[c].dropna().astype(str).head(200)
        longest = max([len(str(c))] + [len(v) for v in sample]) if len(sample) else len(str(c))
        widths.append(min(32, max(9, longest + 2)))
    st.setup_sheet(ws, widths, "4B5563", DOC)
    for j, c in enumerate(export.columns, start=1):
        cell = ws.cell(row=1, column=j, value=str(c))
        cell.font = st.HEADER
        cell.fill = st.fill(st.BRAND)
        cell.alignment = st.WRAP_CENTER
    ws.row_dimensions[1].height = 32
    excluded_font = st.font(11, color=st.GREY_INK)
    has_flag = "INCLUS_ANALYSE" in front
    for i, values in enumerate(export.itertuples(index=False), start=2):
        for j, v in enumerate(values, start=1):
            cell = ws.cell(row=i, column=j, value=_cell_value(v))
            if isinstance(cell.value, datetime.datetime):
                cell.number_format = "dd/mm/yyyy"
        if has_flag and values[0] == 0:
            for j in range(1, len(values) + 1):
                ws.cell(row=i, column=j).font = excluded_font
                ws.cell(row=i, column=j).fill = st.fill(st.GREY_BG)
    ws.freeze_panes = "C2" if front else "B2"
    ws.auto_filter.ref = ws.dimensions

    # --- Dictionnaire des colonnes
    ws = wb.create_sheet(SHEET_DICT)
    st.setup_sheet(ws, [28, 10, 12, 54, 26], "4B5563", DOC)
    row = st.title_block(ws, f"Insight Medics · {DOC}", "Dictionnaire des colonnes",
                         "Ce que contient chaque colonne de la base d'analyse.", 5)
    row = st.header_row(ws, row, ["Colonne", "Type", "Renseignées", "Réponses ou valeurs",
                                  "Origine"])
    corrected = {d["column"] for d in corrections if d.get("kind") != "trim"}
    derived_names = {d.get("name"): d for d in derived}
    labels = {d.get("name"): d.get("label") for d in plan.get("derived") or []}
    for c in export.columns:
        if c == "INCLUS_ANALYSE":
            origin, values = "Ajoutée par nous", "1 = analysé, 0 = écarté (motif à côté)"
        elif c == "MOTIF_EXCLUSION":
            origin, values = "Ajoutée par nous", "Pourquoi la ligne est écartée"
        else:
            values = _value_summary(export[c])
            if c in derived_names:
                kind = KIND_LABEL.get(derived_names[c].get("kind", ""), "Calcul")
                origin = f"Colonne calculée ({kind.lower()})"
                if labels.get(c):
                    values = f"{labels[c]} · {values}"
            elif c in corrected:
                origin = "Votre fichier, corrigée (voir le journal)"
            else:
                origin = "Votre fichier"
        row = st.data_row(ws, row, [str(c), _kind(export[c]), int(export[c].notna().sum()),
                                    values, origin], bold_cols=(1,), center_cols=(2, 3))
    ws.auto_filter.ref = f"A{row - len(export.columns) - 1}:E{row - 1}"

    # --- Journal des corrections
    if corrections:
        ws = wb.create_sheet(SHEET_LOG)
        st.setup_sheet(ws, [20, 22, 40, 10, 24, 30, 34], "B42318", DOC)
        row = st.title_block(ws, f"Insight Medics · {DOC}", "Journal des corrections",
                             "Chaque correction apportée à la base, avec les lignes de votre "
                             "fichier. Votre fichier d'origine n'est pas modifié ; rien n'est "
                             "inventé : une valeur impossible ou illisible devient une case "
                             "vide.", 7)
        row = st.header_row(ws, row, ["Colonne", "Correction", "Motif", "Cellules",
                                      "Lignes du fichier", "Exemples", "Remarques"])
        for d in corrections:
            lines = ", ".join(str(n) for n in d.get("lines") or [])
            if d.get("n", 0) > len(d.get("lines") or []) and d.get("lines"):
                lines += " …"
            remarks = "; ".join(filter(None, [
                f"source : {d['source']}" if d.get("source") else "",
                "illisibles, laissées vides : lignes "
                + ", ".join(str(n) for n in d["illisibles"]) if d.get("illisibles") else "",
                f"{d['censored']} valeur(s) « > » ou « < » laissée(s) vide(s)"
                if d.get("censored") else ""]))
            row = st.data_row(ws, row, [d["column"], CORRECTION_LABEL.get(d["kind"], d["kind"]),
                                        d.get("reason", ""), d.get("n", 0), lines,
                                        "\n".join(d.get("examples") or []), remarks],
                              bold_cols=(1,), center_cols=(4,))

    # --- Colonnes calculées
    if derived:
        ws = wb.create_sheet(SHEET_DERIVED)
        st.setup_sheet(ws, [22, 30, 20, 50, 12, 34], "8A4B08", DOC)
        row = st.title_block(ws, f"Insight Medics · {DOC}", "Colonnes calculées",
                             "Les colonnes que nous avons ajoutées à la base pour l'analyse.", 6)
        row = st.header_row(ws, row, ["Nom", "Libellé", "Calcul", "Pourquoi", "Renseignées",
                                      "Remarques"])
        rationale = {d.get("name"): d.get("rationale", "") for d in plan.get("derived") or []}
        for d in derived:
            remarks = "; ".join(f"{k} : {v}" for k, v in d.items()
                                if k not in {"name", "label", "kind", "n_renseignes", "step"})
            row = st.data_row(ws, row, [d["name"], d.get("label", ""),
                                        KIND_LABEL.get(d.get("kind", ""), d.get("kind", "")),
                                        rationale.get(d["name"], ""), d.get("n_renseignes", ""),
                                        remarks], bold_cols=(1,), center_cols=(5,))

    # --- Sommaire (en premier)
    ws = summary
    st.setup_sheet(ws, [6, 70, 30], st.BRAND, DOC, landscape=False)
    row = st.title_block(ws, f"Insight Medics · {DOC}", "Vos résultats",
                         plan.get("title") or "", 3)
    final = next((f for f in reversed(flow) if "n" in f and f is flow[-1]), None)
    if final:
        row = st.merged_text(ws, row, 1, 3, f"{final['n']} patients analysés sur "
                             f"{flow[0].get('n', '?')} lignes de la base.", st.BODY_BOLD)
    row = st.section(ws, row + 1, "Les tableaux", 3)
    for caption, target in anchors:
        st.internal_link(ws, row, 2, caption, f"'{SHEET_TABLES}'!A{target}")
        st.fit_height(ws, row, [(caption, st.col_width(ws, 2), 11)])
        row += 1
    row = st.section(ws, row + 1, "Les onglets", 3)
    guide = [(SHEET_TABLES, "les tableaux de la partie Résultats, prêts à recopier"),
             (SHEET_FLOW, "combien de patients ont été analysés, et pourquoi"),
             (SHEET_BASE, "la base nettoyée, une ligne par patient ; lignes écartées en gris"),
             (SHEET_DICT, "ce que contient chaque colonne")]
    if corrections:
        guide.append((SHEET_LOG, "chaque correction apportée à votre base"))
    if derived:
        guide.append((SHEET_DERIVED, "les colonnes ajoutées pour l'analyse"))
    for sheet, what in guide:
        st.link(ws, row, 2, sheet, sheet)
        ws.cell(row=row, column=3, value=what).font = st.SMALL
        ws.cell(row=row, column=3).alignment = st.WRAP_TOP
        st.fit_height(ws, row, [(what, st.col_width(ws, 3), 9)])
        row += 1
    st.merged_text(ws, row + 1, 1, 3, "Analyses réalisées par l'équipe Insight Medics, relues "
                   "par un médecin et un biostatisticien.", st.SMALL)
    wb.active = 0

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
