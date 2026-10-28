"""Rendu Word d'un livrable de thèse selon la charte de la faculté (python-docx).

Une section par page intercalaire (sans en-tête ni numéro affiché, page comptée), puis une
section par chapitre avec son en-tête et son pied de page. Les titres utilisent les
styles Word « Heading 1 » à « Heading 4 » (« Titre 1 »… dans un Word français), pour
que la table des matières de l'étudiant les reprenne ; les légendes utilisent le style
« Caption » (« Légende »), pour ses listes des tableaux et des figures.

Les polices des gabarits (AvantGarde, Albertus, Souvenir, Lucian BT…) sont déclarées
dans les styles même si elles manquent sur le serveur : Word les affiche chez
l'étudiant qui les a, et les remplace sinon (comme sur le poste de Hamza).
"""

from __future__ import annotations

import io
from collections.abc import Iterable
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .charte import Align, Charte, Font, Heading, Para, Rule, TableStyle
from .document import (
    BOLD_RE,
    CALL_RE,
    BulletsBlock,
    Chapter,
    FigureBlock,
    HeadingBlock,
    ParagraphBlock,
    ReferencesBlock,
    TableBlock,
    ThesisDoc,
)
from .typo import NBSP, en_typo, fr_typo, roman

_ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
          "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}
HEADING_STYLES = ("Heading 1", "Heading 2", "Heading 3", "Heading 4")
TWIPS_PER_CM = 567


class RenderError(ValueError):
    """Contenu impossible à mettre en page (appel sans tableau, image absente…)."""


# ---------------------------------------------------------------- mise en forme bas niveau

def _font_xml(rpr: Any, font: Font, lang: str | None = None) -> None:
    """Police complète d'un `w:rPr` (style ou run) : sans police de thème, sinon Word
    garde la police du thème à la place de celle de la charte."""
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn(f"w:{attr}"), None)
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        fonts.set(qn(f"w:{attr}"), font.name)
    if lang:
        el = rpr.find(qn("w:lang"))
        if el is None:
            el = OxmlElement("w:lang")
            rpr.append(el)
        el.set(qn("w:val"), lang)


def _apply_font(target: Any, font: Font, lang: str | None = None) -> None:
    """`target` : un style ou un run python-docx."""
    f = target.font
    f.size = Pt(font.size)
    f.bold = font.bold
    f.italic = font.italic
    f.small_caps = font.small_caps
    f.color.rgb = RGBColor.from_string(font.color) if font.color else None
    rpr = target.element.get_or_add_rPr()
    _font_xml(rpr, font, lang)


def _apply_para(fmt: Any, para: Para) -> None:
    fmt.alignment = _ALIGN[para.align]
    fmt.space_before = Pt(para.before)
    fmt.space_after = Pt(para.after)
    fmt.line_spacing = para.line
    fmt.first_line_indent = Cm(para.first_indent_cm) if para.first_indent_cm else Cm(0)
    fmt.left_indent = Cm(para.left_indent_cm) if para.left_indent_cm else Cm(0)
    fmt.keep_with_next = para.keep_next
    fmt.widow_control = True


def _border_el(tag: str, rule: Rule) -> Any:
    el = OxmlElement(f"w:{tag}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(max(2, round(rule.pt * 8))))   # huitièmes de point
    el.set(qn("w:space"), "1")
    el.set(qn("w:color"), rule.color or "auto")
    return el


def _para_border(paragraph: Any, side: str, rule: Rule) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    bdr = ppr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        ppr.append(bdr)
    bdr.append(_border_el(side, rule))


def _style_border(style: Any, side: str, rule: Rule) -> None:
    ppr = style.element.get_or_add_pPr()
    bdr = ppr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        ppr.append(bdr)
    bdr.append(_border_el(side, rule))


def _field(paragraph: Any, instr: str, font: Font | None = None, placeholder: str = "1") -> None:
    """Champ Word (PAGE, TOC…) : begin / instr / separate / valeur / end."""
    def run_with(child: Any) -> None:
        r = paragraph.add_run()
        if font:
            _apply_font(r, font)
        r._r.append(child)

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run_with(begin)
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = f" {instr} "
    run_with(text)
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run_with(sep)
    r = paragraph.add_run(placeholder)
    if font:
        _apply_font(r, font)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run_with(end)


# ---------------------------------------------------------------- document et styles

class _Ctx:
    def __init__(self, doc: Any, charte: Charte, thesis: ThesisDoc):
        self.doc = doc
        self.c = charte
        self.thesis = thesis
        self.tables: dict[str, str] = {}      # clé → « Tableau III »
        self.figures: dict[str, str] = {}     # clé → « Figure 2 »
        self.next_table = thesis.first_table
        self.next_figure = thesis.first_figure
        self.counters = [0, 0, 0, 0]
        self.bullet_num_id: int | None = None


def _setup_styles(doc: Any, c: Charte) -> None:
    normal = doc.styles["Normal"]
    _apply_font(normal, c.body, c.lang)
    _apply_para(normal.paragraph_format, c.body_para)
    for i, style_name in enumerate(HEADING_STYLES):
        style = doc.styles[style_name]
        level = c.levels[min(i, len(c.levels) - 1)]
        _apply_font(style, level.font, c.lang_small_caps if level.font.small_caps else c.lang)
        _apply_para(style.paragraph_format, level.para)
        # Le style intégré « Heading » a une couleur de thème : la charte la remplace.
        style.font.color.rgb = (RGBColor.from_string(level.font.color) if level.font.color
                                else RGBColor(0, 0, 0))
    cap = doc.styles["Caption"]
    _apply_font(cap, c.caption_table.font, c.lang)
    _apply_para(cap.paragraph_format, c.caption_table.para)


def _setup_section(section: Any, c: Charte) -> None:
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    top, bottom, left, right = c.margins_cm
    section.top_margin, section.bottom_margin = Cm(top), Cm(bottom)
    section.left_margin, section.right_margin = Cm(left), Cm(right)
    section.header_distance, section.footer_distance = Cm(c.header_cm), Cm(c.footer_cm)


def _page_start(section: Any, start: int | None) -> None:
    sect = section._sectPr
    pg = sect.find(qn("w:pgNumType"))
    if start is None:
        if pg is not None:
            sect.remove(pg)
        return
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        sect.append(pg)
    pg.set(qn("w:start"), str(start))


def _clear(part: Any) -> None:
    for p in list(part.paragraphs):
        p._p.getparent().remove(p._p)
    for t in list(part.tables):
        t._tbl.getparent().remove(t._tbl)


def _header_footer(section: Any, ctx: _Ctx, chapter_name: str, empty: bool) -> None:
    c = ctx.c
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    _clear(section.header)
    _clear(section.footer)
    head = section.header.add_paragraph()
    foot = section.footer.add_paragraph()
    if empty:
        return
    h = c.header
    if h.kind == "chapter":
        head.alignment = _ALIGN[h.align]
        _apply_font(head.add_run(chapter_name), h.font)
    else:
        # Titre de la thèse à gauche, chapitre à droite (taquet à droite).
        width = Cm(c.usable_cm)
        head.paragraph_format.tab_stops.add_tab_stop(width, alignment=2)  # 2 = droite
        title = ctx.thesis.thesis_title or ""
        if len(title) > 70:
            title = title[:67].rstrip() + "…"
        _apply_font(head.add_run(title), h.title_font or h.font)
        _apply_font(head.add_run("\t" + chapter_name), h.font)
    if h.rule_below:
        _para_border(head, "bottom", h.rule_below)
    f = c.footer
    foot.alignment = _ALIGN[f.align]
    _field(foot, "PAGE", f.font)
    if f.rule_above:
        _para_border(foot, "top", f.rule_above)


# ---------------------------------------------------------------- texte courant

def _resolve_calls(text: str, ctx: _Ctx) -> str:
    def repl(m: Any) -> str:
        kind, key = m.group(1), m.group(2)
        table = ctx.tables if kind == "TAB" else ctx.figures
        if key not in table:
            raise RenderError(f"appel {{{{{kind}:{key}}}}} sans {'tableau' if kind == 'TAB' else 'figure'} "
                              "correspondant dans le chapitre")
        return table[key]
    return CALL_RE.sub(repl, text)


def _typo(text: str, ctx: _Ctx) -> str:
    """Typographie française ; en anglais, seulement les intervalles et l'apostrophe."""
    if ctx.thesis.lang == "fr":
        return fr_typo(text)
    return en_typo(text)


def _add_runs(paragraph: Any, text: str, font: Font | None = None,
              ctx: _Ctx | None = None) -> None:
    """Texte avec `**gras**` ; typographie de la langue appliquée."""
    text = _typo(text, ctx) if ctx else fr_typo(text)
    pos = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            r = paragraph.add_run(text[pos:m.start()])
            if font:
                _apply_font(r, font)
        r = paragraph.add_run(m.group(1))
        if font:
            _apply_font(r, font)
        r.bold = True
        pos = m.end()
    if pos < len(text):
        r = paragraph.add_run(text[pos:])
        if font:
            _apply_font(r, font)


def _heading_label(level: int, ctx: _Ctx, number: int | None = None) -> str:
    """Numéro du titre : FMS « I. », « 1. », « 1.1. » ; autres « 1. », « 1.1. »."""
    ctx.counters[level] = number if number is not None else ctx.counters[level] + 1
    for i in range(level + 1, len(ctx.counters)):
        ctx.counters[i] = 0
    if ctx.c.roman_top_level:
        if level == 0:
            return roman(ctx.counters[0]) + "."
        return ".".join(str(n) for n in ctx.counters[1:level + 1]) + "."
    return ".".join(str(n) for n in ctx.counters[:level + 1]) + "."


def _heading(block: HeadingBlock, ctx: _Ctx) -> None:
    c = ctx.c
    level = min(block.level, len(c.levels) - 1)
    spec = c.levels[level]
    text = block.text.strip().rstrip(":").strip()
    if spec.font.caps:
        text = text.upper()
    label = _heading_label(level, ctx, block.number)
    p = ctx.doc.add_paragraph(style=HEADING_STYLES[min(level, 3)])
    suffix = spec.suffix.replace(" ", NBSP) if ctx.thesis.lang == "fr" else spec.suffix.strip()
    run = p.add_run(f"{label} {_typo(text, ctx)}{suffix}")
    if spec.font.small_caps and c.lang_small_caps:
        _font_xml(run._r.get_or_add_rPr(), spec.font, c.lang_small_caps)


def _paragraph(text: str, ctx: _Ctx) -> None:
    p = ctx.doc.add_paragraph()
    _add_runs(p, _resolve_calls(text, ctx), ctx=ctx)


def _bullet_numbering(ctx: _Ctx) -> int:
    """Définition de liste à puces de la charte (symbole, retraits), créée une fois."""
    if ctx.bullet_num_id is not None:
        return ctx.bullet_num_id
    numbering = ctx.doc.part.numbering_part.element
    b = ctx.c.bullets
    abstract_ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    for tag, val in (("w:start", "1"), ("w:numFmt", "bullet"), ("w:lvlText", b.symbol),
                     ("w:lvlJc", "left")):
        el = OxmlElement(tag)
        el.set(qn("w:val"), val)
        lvl.append(el)
    ppr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(round(b.left_cm * TWIPS_PER_CM)))
    ind.set(qn("w:hanging"), str(round(b.hanging_cm * TWIPS_PER_CM)))
    ppr.append(ind)
    lvl.append(ppr)
    abstract.append(lvl)
    # Les abstractNum précèdent les num dans numbering.xml.
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    numbering.append(num)
    ctx.bullet_num_id = num_id
    return num_id


def _bullets(block: BulletsBlock, ctx: _Ctx) -> None:
    num_id = _bullet_numbering(ctx)
    for i, item in enumerate(block.items):
        p = ctx.doc.add_paragraph()
        ppr = p._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        nid = OxmlElement("w:numId")
        nid.set(qn("w:val"), str(num_id))
        num_pr.append(ilvl)
        num_pr.append(nid)
        ppr.append(num_pr)
        fmt = p.paragraph_format
        fmt.first_line_indent = Cm(-ctx.c.bullets.hanging_cm)
        fmt.left_indent = Cm(ctx.c.bullets.left_cm)
        fmt.space_after = Pt(ctx.c.bullets.after if i < len(block.items) - 1
                             else ctx.c.body_para.after)
        _add_runs(p, _resolve_calls(item, ctx), ctx=ctx)


# ---------------------------------------------------------------- légendes, notes

def _word(key: str, ctx: _Ctx) -> str:
    return {"table": ("Tableau", "Table")}[key][0 if ctx.thesis.lang == "fr" else 1]


def _caption(label: str, title: str, spec: Any, ctx: _Ctx, keep_next: bool) -> None:
    p = ctx.doc.add_paragraph(style="Caption")
    _apply_para(p.paragraph_format, spec.para)
    p.paragraph_format.keep_with_next = keep_next
    title = _typo(title.strip().rstrip("."), ctx)
    lab = p.add_run(f"{label}{NBSP}: " if ctx.thesis.lang == "fr" else f"{label}: ")
    _apply_font(lab, spec.font)
    if spec.label_bold_only:
        lab.bold = True
    body = p.add_run(title)
    _apply_font(body, spec.font)
    if spec.label_bold_only:
        body.bold = False


def _note(text: str, ctx: _Ctx, align: Align | None = None) -> None:
    if not text.strip():
        return
    p = ctx.doc.add_paragraph()
    para = ctx.c.note_para
    _apply_para(p.paragraph_format, Para(align or para.align, para.before, para.after, para.line))
    _add_runs(p, _resolve_calls(text, ctx), ctx.c.note, ctx)


# ---------------------------------------------------------------- tableaux

def _cell_border(cell: Any, side: str, rule: Rule) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    borders.append(_border_el(side, rule))


def _shade(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _row_flag(row: Any, tag: str) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement(f"w:{tag}")
    tr_pr.append(el)


def _widths(block: TableBlock, total_cm: float) -> list[float]:
    n = len(block.columns)
    if block.col_widths and len(block.col_widths) == n:
        weights = block.col_widths
    elif n == 1:
        weights = [1.0]
    else:
        weights = [2.2] + [1.0] * (n - 1)
        if n > 5:
            weights = [1.8] + [1.0] * (n - 1)
    s = sum(weights)
    return [total_cm * w / s for w in weights]


def _is_group(row: list[str], i: int, block: TableBlock) -> bool:
    if block.group_rows is not None:
        return i in block.group_rows
    return len(row) > 1 and bool(str(row[0]).strip()) and all(not str(v).strip() for v in row[1:])


def _table(block: TableBlock, ctx: _Ctx) -> None:
    c: Charte = ctx.c
    ts: TableStyle = c.table
    label = f"{_word('table', ctx)}{NBSP}{roman(ctx.next_table)}"
    ctx.tables[block.key] = label
    ctx.next_table += 1
    _caption(label, block.title, c.caption_table, ctx, keep_next=True)

    n = len(block.columns)
    table = ctx.doc.add_table(rows=1 + len(block.rows), cols=n)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(round(ts.width_cm * TWIPS_PER_CM)))
    widths = _widths(block, ts.width_cm)

    def fill(cell: Any, text: str, j: int, bold: bool = False, indent: bool = False,
             color: str | None = None) -> None:
        cell.width = Cm(widths[j])
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(ts.cell_before)
        p.paragraph_format.space_after = Pt(ts.cell_after)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.first_line_indent = Cm(0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
        if indent:
            p.paragraph_format.left_indent = Cm(ts.indent_cm)
        font = Font(ts.font.name, ts.font.size, bold=bold, color=color or ts.text_color)
        _add_runs(p, text, font, ctx)

    header = table.rows[0]
    _row_flag(header, "tblHeader")
    _row_flag(header, "cantSplit")
    for j, title in enumerate(block.columns):
        cell = header.cells[j]
        fill(cell, str(title), j, bold=True)
        _cell_border(cell, "top", ts.rule_top)
        _cell_border(cell, "bottom", ts.rule_header)
    band = 0
    for i, row in enumerate(block.rows):
        tr = table.rows[i + 1]
        _row_flag(tr, "cantSplit")
        group = _is_group(row, i, block)
        if not group:
            band += 1
        values = list(row) + [""] * (n - len(row))
        for j, value in enumerate(values[:n]):
            text = str(value)
            indent = j == 0 and text.startswith(" ")
            cell = tr.cells[j]
            fill(cell, text.strip(), j, bold=group, indent=indent,
                 color=ts.group_color if group else None)
            if group and ts.group_fill:
                _shade(cell, ts.group_fill)
            elif not group and ts.band_fill and band % 2 == 0:
                _shade(cell, ts.band_fill)
            if i == len(block.rows) - 1:
                _cell_border(cell, "bottom", ts.rule_bottom)
            # Les deux dernières lignes restent avec la note.
            if i >= len(block.rows) - 2 and block.note:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = True
        # L'en-tête reste lié aux deux premières lignes.
    for tr in table.rows[:3]:
        for cell in tr.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = True
    _note(block.note, ctx)
    if not block.note:
        spacer = ctx.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(0)
        spacer.paragraph_format.line_spacing = 1.0


# ---------------------------------------------------------------- figures

def _figure(block: FigureBlock, ctx: _Ctx) -> None:
    c = ctx.c
    path = Path(block.image_path)
    if not block.image_path or not path.is_file():
        raise RenderError(f"figure « {block.key} » : image introuvable")
    label = f"Figure{NBSP}{ctx.next_figure}"
    ctx.figures[block.key] = label  # « Figure » dans les deux langues
    ctx.next_figure += 1
    p = ctx.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(str(path), width=Cm(min(block.width_cm or c.usable_cm, c.usable_cm)))
    _caption(label, block.title, c.caption_figure, ctx, keep_next=bool(block.note))
    _note(block.note, ctx, align="center" if c.code == "FMSf" else None)


def _references(block: ReferencesBlock, ctx: _Ctx) -> None:
    c = ctx.c
    heading = ctx.doc.add_paragraph(style=HEADING_STYLES[0])
    title = block.title.strip().rstrip(":").strip().upper()
    heading.add_run(title + c.levels[0].suffix.replace(" ", NBSP))
    for k, item in enumerate(block.items, start=block.start):
        p = ctx.doc.add_paragraph()
        _apply_para(p.paragraph_format, c.refs.para)
        p.paragraph_format.left_indent = Cm(c.refs.hanging_cm)
        p.paragraph_format.first_line_indent = Cm(-c.refs.hanging_cm)
        _add_runs(p, f"{k}. {item.strip()}", c.refs.font, ctx)


# ---------------------------------------------------------------- chapitres

def _intercalaire(chapter: Chapter, ctx: _Ctx) -> None:
    inter = ctx.c.intercalaire
    assert inter is not None
    for _ in range(inter.blank_paragraphs):
        ctx.doc.add_paragraph()
    p = ctx.doc.add_paragraph()
    spec: Heading = inter.heading
    _apply_para(p.paragraph_format, spec.para)
    p.paragraph_format.first_line_indent = Cm(0)
    if inter.vertical_center:
        p.paragraph_format.space_before = Pt(260)
    text = chapter.name.upper() if inter.uppercase else chapter.name
    _apply_font(p.add_run(text), spec.font,
                ctx.c.lang_small_caps if spec.font.small_caps else None)
    if spec.rule_below:
        _para_border(p, "bottom", spec.rule_below)


def _chapter_title(chapter: Chapter, ctx: _Ctx) -> None:
    spec = ctx.c.chapter_title
    if spec is None:
        return
    p = ctx.doc.add_paragraph()
    _apply_para(p.paragraph_format, spec.para)
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_with_next = True
    text = chapter.name.upper() if ctx.c.chapter_title_uppercase else chapter.name
    _apply_font(p.add_run(_typo(text, ctx)), spec.font)
    if spec.rule_below:
        _para_border(p, "bottom", spec.rule_below)


def _blocks(chapter: Chapter, ctx: _Ctx) -> None:
    # Les appels se résolvent sur les tableaux et figures du chapitre : numérotation
    # préalable dans l'ordre d'apparition (un appel peut précéder son tableau).
    t, f = ctx.next_table, ctx.next_figure
    for b in chapter.blocks:
        if isinstance(b, TableBlock):
            ctx.tables.setdefault(b.key, f"{_word('table', ctx)}{NBSP}{roman(t)}")
            t += 1
        elif isinstance(b, FigureBlock):
            ctx.figures.setdefault(b.key, f"Figure{NBSP}{f}")
            f += 1
    ctx.counters = [0, 0, 0, 0]
    skip_headings = chapter.key in ctx.c.no_headings_in
    for b in chapter.blocks:
        if isinstance(b, HeadingBlock):
            if not skip_headings:
                _heading(b, ctx)
        elif isinstance(b, ParagraphBlock):
            _paragraph(b.text, ctx)
        elif isinstance(b, BulletsBlock):
            _bullets(b, ctx)
        elif isinstance(b, TableBlock):
            _table(b, ctx)
        elif isinstance(b, FigureBlock):
            _figure(b, ctx)
        elif isinstance(b, ReferencesBlock):
            _references(b, ctx)


def render(thesis: ThesisDoc, charte: Charte, title: str = "") -> bytes:
    """Met en page le livrable ; renvoie le .docx. Lève `RenderError` si le contenu est
    incohérent (appel sans tableau, image manquante)."""
    doc = Document()
    c = charte
    _setup_styles(doc, c)
    core = doc.core_properties
    core.author = ""
    core.last_modified_by = ""
    core.comments = ""
    core.title = title
    ctx = _Ctx(doc, c, thesis)
    first = True
    page = thesis.first_page
    for chapter in thesis.chapters:
        if thesis.insert:
            section = doc.sections[0] if first else doc.add_section(WD_SECTION.NEW_PAGE)
            _setup_section(section, c)
            _header_footer(section, ctx, chapter.name, empty=False)
            _blocks(chapter, ctx)
            first = False
            continue
        if c.intercalaire is not None:
            section = doc.sections[0] if first else doc.add_section(WD_SECTION.NEW_PAGE)
            _setup_section(section, c)
            _page_start(section, page if first else None)
            _header_footer(section, ctx, chapter.name, empty=True)
            _intercalaire(chapter, ctx)
            section = doc.add_section(WD_SECTION.NEW_PAGE)
            _page_start(section, None)
        else:
            section = doc.sections[0] if first else doc.add_section(WD_SECTION.NEW_PAGE)
            _page_start(section, page if first else None)
        _setup_section(section, c)
        _header_footer(section, ctx, chapter.name, empty=False)
        _chapter_title(chapter, ctx)
        _blocks(chapter, ctx)
        first = False
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def plain_text(doc_bytes: bytes) -> Iterable[str]:
    """Paragraphes du Word (corps et tableaux), pour les contrôles et les tests."""
    doc = Document(io.BytesIO(doc_bytes))
    for p in doc.paragraphs:
        yield p.text
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                yield cell.text

