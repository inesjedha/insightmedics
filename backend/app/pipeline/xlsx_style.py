"""Socle commun des classeurs Excel d'Insight Medics (audit étudiant, résultats, équipe).

Un seul endroit pour l'apparence : palette du site, une police, des styles nommés, des
hauteurs de ligne calculées d'après le texte (Excel n'agrandit jamais une cellule
fusionnée, et openpyxl n'écrit aucune hauteur), les nombres à la française, et une mise
en page d'impression A4 paysage avec pied de page.
"""

from __future__ import annotations

import math
from collections.abc import Sequence
from typing import Any

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.worksheet import Worksheet

# ---------------------------------------------------------------- palette

FONT = "Calibri"
BRAND = "1F3FD1"        # bleu Insight Medics (site)
BRAND_DARK = "16289A"
INK = "111827"
MUTED_INK = "4B5563"
LINE_GREY = "E5E7EB"
SURFACE = "F3F5FB"      # fond des tuiles et intertitres neutres
INPUT = "FFF8D6"        # cases à remplir par l'étudiant
RED_BG, RED_INK = "FBE4E1", "9F1C12"
ORANGE_BG, ORANGE_INK = "FDF0D8", "8A4B08"
GREEN_BG, GREEN_INK = "E3F4EA", "1E6B3A"
GREY_BG, GREY_INK = "EEF0F4", "6B7280"


def font(size: float = 11, bold: bool = False, color: str = INK, italic: bool = False,
         underline: str | None = None) -> Font:
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic,
                underline=underline)


def fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


TITLE = font(18, bold=True)
KICKER = font(9, color=MUTED_INK)
SUBTITLE = font(11, color=MUTED_INK)
SECTION = font(13, bold=True, color=BRAND_DARK)
BODY = font(11)
BODY_BOLD = font(11, bold=True)
SMALL = font(9, color=MUTED_INK)
HEADER = font(10.5, bold=True, color="FFFFFF")
HEADER_INPUT = font(10.5, bold=True, color=INK)
LINK = font(11, color=BRAND, underline="single")

WRAP_TOP = Alignment(wrap_text=True, vertical="top")
WRAP_CENTER = Alignment(wrap_text=True, vertical="center")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)

THIN = Side(style="thin", color=LINE_GREY)
ROW_LINE = Border(bottom=THIN)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
SECTION_LINE = Border(bottom=Side(style="medium", color=BRAND))

# Feu des missions et priorités : fond, encre, libellé.
LIGHT = {"rouge": (RED_BG, RED_INK, "Rouge"), "orange": (ORANGE_BG, ORANGE_INK, "Orange"),
         "vert": (GREEN_BG, GREEN_INK, "Vert")}
PRIORITY_STYLE = {"avant_analyse": (RED_BG, RED_INK), "importante": (ORANGE_BG, ORANGE_INK),
                  "facultative": (GREY_BG, GREY_INK)}

LINE_HEIGHT = 15.0      # points par ligne de texte en Calibri 11
CHARS_PER_UNIT = 1.15   # caractères par unité de largeur de colonne (texte courant)


# ---------------------------------------------------------------- nombres

def fr_num(value: Any, decimals: int | None = None) -> str:
    """Nombre à la française : virgule décimale, espace fine pour les milliers."""
    if value is None or value == "":
        return "—"
    if isinstance(value, (int, float)):
        if decimals is not None:
            text = f"{value:,.{decimals}f}"
        elif float(value).is_integer():
            text = f"{int(value):,}"
        else:
            text = f"{value:,.1f}"
        return text.replace(",", " ").replace(".", ",")
    return str(value).replace(".", ",")


def fr_pct(value: Any, decimals: int = 1) -> str:
    """« 86,7 % » (valeur déjà en pourcentage)."""
    if value is None or value == "":
        return "—"
    return f"{fr_num(float(value), decimals)} %"


# ---------------------------------------------------------------- hauteurs de ligne

def text_lines(text: Any, width_units: float, size: float = 11) -> int:
    """Nombre de lignes qu'occupe `text` dans une largeur de `width_units` (estimation
    prudente : mieux vaut une ligne de trop qu'une phrase coupée)."""
    if text is None or text == "":
        return 1
    per_line = max(4.0, width_units * CHARS_PER_UNIT * 11 / size)
    return sum(max(1, math.ceil(len(part) / per_line)) for part in str(text).split("\n"))


def col_width(ws: Worksheet, col: int) -> float:
    dim = ws.column_dimensions.get(get_column_letter(col))
    return float(dim.width) if dim is not None and dim.width else 8.43


def span_width(ws: Worksheet, first: int, last: int) -> float:
    return sum(col_width(ws, c) for c in range(first, last + 1))


def fit_height(ws: Worksheet, row: int, cells: Sequence[tuple[Any, float, float]],
               minimum: float = LINE_HEIGHT + 3) -> None:
    """Hauteur de la ligne d'après ses textes : [(texte, largeur, taille de police)]."""
    lines = max((text_lines(t, w, s) * (LINE_HEIGHT * s / 11) for t, w, s in cells),
                default=LINE_HEIGHT)
    ws.row_dimensions[row].height = max(minimum, lines + 5)


# ---------------------------------------------------------------- feuille

def setup_sheet(ws: Worksheet, widths: Sequence[float], tab_color: str, document: str,
                landscape: bool = True) -> None:
    """Largeurs, quadrillage masqué, couleur d'onglet, impression A4 ajustée à la largeur,
    pied de page « Insight Medics · document · page x sur y »."""
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100
    ws.sheet_properties.tabColor = tab_color
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ps = ws.page_setup
    ps.paperSize = ws.PAPERSIZE_A4
    ps.orientation = "landscape" if landscape else "portrait"
    ps.fitToWidth = 1
    ps.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.6
    ws.oddFooter.left.text = f"Insight Medics · {document}"
    ws.oddFooter.left.size = 8
    ws.oddFooter.right.text = "Page &P sur &N"
    ws.oddFooter.right.size = 8


def merged_text(ws: Worksheet, row: int, first: int, last: int, value: Any,
                style: Font = BODY, fill_color: str | None = None,
                alignment: Alignment = WRAP_TOP, min_height: float | None = None) -> int:
    """Texte fusionné sur les colonnes first..last, hauteur calculée. Renvoie row + 1."""
    c = ws.cell(row=row, column=first, value=value)
    c.font = style
    c.alignment = alignment
    if fill_color:
        for j in range(first, last + 1):
            ws.cell(row=row, column=j).fill = fill(fill_color)
    if last > first:
        ws.merge_cells(start_row=row, start_column=first, end_row=row, end_column=last)
    size = style.sz or 11
    fit_height(ws, row, [(value, span_width(ws, first, last), size)],
               minimum=min_height or LINE_HEIGHT * size / 11 + 4)
    return row + 1


def title_block(ws: Worksheet, kicker: str, title: str, subtitle: str, last_col: int) -> int:
    """Surtitre, titre, sous-titre ; renvoie la première ligne libre (après un blanc)."""
    row = merged_text(ws, 1, 1, last_col, kicker, KICKER)
    row = merged_text(ws, row, 1, last_col, title, TITLE, min_height=30)
    if subtitle:
        row = merged_text(ws, row, 1, last_col, subtitle, SUBTITLE)
    return row + 1


def section(ws: Worksheet, row: int, text: str, last_col: int) -> int:
    """Intertitre souligné du bleu de la marque ; renvoie la ligne suivante."""
    for j in range(1, last_col + 1):
        ws.cell(row=row, column=j).border = SECTION_LINE
    c = ws.cell(row=row, column=1, value=text)
    c.font = SECTION
    c.alignment = Alignment(vertical="bottom")
    ws.row_dimensions[row].height = 24
    return row + 1


def header_row(ws: Worksheet, row: int, headers: list[str],
               input_cols: tuple[int, ...] = (), freeze: bool = True) -> int:
    """En-têtes de tableau (bleu ; jaune pour les colonnes à remplir)."""
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=j, value=h)
        is_input = j in input_cols
        c.font = HEADER_INPUT if is_input else HEADER
        c.fill = fill(INPUT if is_input else BRAND)
        c.alignment = WRAP_CENTER
    fit_height(ws, row, [(h, col_width(ws, j), 10.5) for j, h in enumerate(headers, 1)],
               minimum=22)
    if freeze:
        ws.freeze_panes = ws.cell(row=row + 1, column=1)
        ws.print_title_rows = f"{row}:{row}"
    return row + 1


def data_row(ws: Worksheet, row: int, values: list[Any], input_cols: tuple[int, ...] = (),
             bold_cols: tuple[int, ...] = (), center_cols: tuple[int, ...] = ()) -> int:
    """Ligne de tableau : texte replié, filet bas, hauteur calculée."""
    sizes = []
    for j, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=j, value=v)
        c.font = BODY_BOLD if j in bold_cols else BODY
        c.alignment = CENTER if j in center_cols else WRAP_TOP
        c.border = ROW_LINE
        if j in input_cols:
            c.fill = fill(INPUT)
        sizes.append((v, col_width(ws, j), 11))
    fit_height(ws, row, sizes)
    return row + 1


def band(ws: Worksheet, row: int, text: str, bg: str, ink: str, last_col: int,
         size: float = 11.5) -> int:
    """Bandeau d'intertitre coloré (mission, question) sur toute la largeur."""
    return merged_text(ws, row, 1, last_col, text, font(size, bold=True, color=ink), bg,
                       WRAP_CENTER, min_height=24)


def link(ws: Worksheet, row: int, col: int, text: str, sheet: str) -> None:
    """Lien vers un autre onglet du classeur."""
    internal_link(ws, row, col, text, f"'{sheet}'!A1")


def internal_link(ws: Worksheet, row: int, col: int, text: str, location: str) -> None:
    """Lien interne au classeur (« 'Tableaux'!A12 »), suivi par Excel et LibreOffice."""
    c = ws.cell(row=row, column=col, value=text)
    c.hyperlink = Hyperlink(ref=c.coordinate, location=location, display=text)
    c.font = LINK
    c.alignment = WRAP_TOP


def autofit_rows(ws: Worksheet, first_row: int = 1) -> None:
    """Hauteur de chaque ligne d'après ses cellules à texte replié (hors cellules
    fusionnées), pour les feuilles écrites sans `data_row`."""
    merged = {c for r in ws.merged_cells.ranges for row in ws.iter_rows(
        min_row=r.min_row, max_row=r.max_row, min_col=r.min_col, max_col=r.max_col)
        for c in row}
    for row in ws.iter_rows(min_row=first_row):
        cells = [(c.value, col_width(ws, c.column), float(c.font.sz or 11))
                 for c in row if c.value is not None and c not in merged
                 and c.alignment is not None and c.alignment.wrap_text]
        if cells:
            fit_height(ws, row[0].row, cells)


def restyle(ws: Worksheet) -> None:
    """Remplace les polices d'un classeur ancien par celles du socle (même taille, gras,
    couleur ; police Calibri)."""
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and c.font is not None and c.font.name != FONT:
                f = c.font
                c.font = Font(name=FONT, size=f.sz or 11, bold=f.b, italic=f.i,
                              color=f.color)
